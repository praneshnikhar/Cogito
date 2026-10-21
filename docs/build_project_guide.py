#!/usr/bin/env python3
"""Build docs/Cogito_Project_Guide.docx.

A beginner-friendly guide to the project, MongoDB, Redis, Next.js, FastAPI, and
RAG. Run with the system Python (python-docx + matplotlib are installed):

    python3 docs/build_project_guide.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "docs" / "guide_assets"
OUT = ROOT / "docs" / "Cogito_Project_Guide.docx"

NAVY = RGBColor(0x14, 0x20, 0x33)
BLUE = RGBColor(0x25, 0x63, 0xEB)
PURPLE = RGBColor(0x6D, 0x4A, 0xFF)
GREEN = RGBColor(0x0F, 0x9D, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)
MUTED = RGBColor(0x5B, 0x6B, 0x82)
DARK = RGBColor(0x1F, 0x29, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE = "EEF4FF"
LIGHT_GREEN = "E7F7F0"
LIGHT_AMBER = "FFF6E5"
LIGHT_GRAY = "F3F5F9"
LIGHT_LINE = "D7DEEA"
HEADER_FILL = "0B1020"

doc = Document()

# ---------------------------------------------------------------- page + styles
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.14

for name, size, color, before, after in [
    ("Heading 1", 20, NAVY, 18, 8),
    ("Heading 2", 14, BLUE, 14, 4),
    ("Heading 3", 11.5, NAVY, 10, 3),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.1

# ---------------------------------------------------------------- helpers
def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(cell, top=0.06, bottom=0.06, left=0.08, right=0.08) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(int(value * 1440)))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def table_borders(table, color: str = LIGHT_LINE, size: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def H1(text: str) -> None:
    doc.add_heading(text, level=1)


def H2(text: str) -> None:
    doc.add_heading(text, level=2)


def H3(text: str) -> None:
    doc.add_heading(text, level=3)


def P(text: str = "", bold_lead: str | None = None, size: float | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = NAVY
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    return p


def B(items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def N(items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def CODE(text: str, caption: str | None = None, size: float = 8.5) -> None:
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(2)
        r = cp.add_run(caption)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table, "C9D3E3", 4)
    cell = table.cell(0, 0)
    shade(cell, LIGHT_GRAY)
    cell_margins(cell, 0.07, 0.07, 0.1, 0.1)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(size)
        run.font.color.rgb = DARK
        if i < len(lines) - 1:
            p.add_run().add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def CALLOUT(title: str, text: str, color: str = "2563EB", fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.09)
    table.columns[1].width = Inches(6.4)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    shade(left, color)
    shade(right, fill)
    cell_margins(left, 0.02, 0.02, 0.02, 0.02)
    cell_margins(right, 0.08, 0.08, 0.14, 0.12)
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(color)
    p2 = right.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def TABLE(headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 8.8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        shade(cell, HEADER_FILL)
        cell_margins(cell, 0.05, 0.05, 0.07, 0.05)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            shade(cell, "F7F9FC" if ri % 2 == 0 else "FFFFFF")
            cell_margins(cell, 0.05, 0.05, 0.07, 0.05)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = NAVY if ci == 0 else DARK
            r.bold = ci == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(name: str, caption: str, width: float = 6.5) -> None:
    path = ASSETS / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED


def PB() -> None:
    doc.add_page_break()


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_toc() -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Right-click and choose Update Field to build the table of contents."
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(t)
    run._r.append(fld3)


# ---------------------------------------------------------------- header/footer
header = section.header
hp = header.paragraphs[0]
hp.text = "Cogito · Complete Project Guide"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Cogito project guide  ·  page ")
r.font.size = Pt(8)
r.font.color.rgb = MUTED
add_page_field(fp)

# ================================================================== cover
cover = doc.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
cover.columns[0].width = Inches(6.5)
cell = cover.cell(0, 0)
cell.width = Inches(6.5)
shade(cell, HEADER_FILL)
cell_margins(cell, 0.35, 0.35, 0.35, 0.3)

p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(10)
r = p.add_run("COGITO")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x93, 0xA0, 0xBF)

p = cell.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("The Complete\nProject Guide")
r.bold = True
r.font.size = Pt(32)
r.font.color.rgb = WHITE
r.font.name = "Calibri"

p = cell.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("MongoDB  ·  FastAPI  ·  Next.js  ·  Redis")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor(0xC9, 0xD4, 0xE6)

doc.add_paragraph().paragraph_format.space_after = Pt(18)
P("A beginner-friendly explanation of what the project does, why each technology is there, and how the pieces talk to each other.",
  size=13)
P("Prepared for the NoSQL course project and viva.", size=10.5)

doc.add_paragraph()
P("Group", bold_lead="")
for name, roll in [
    ("Pranesh Nikhar", "2427010267"),
    ("Soumyajit Chatterjee", "2427010255"),
    ("Anmol Kaushik", "2427010576"),
    ("Ojas Khandelwal", "2427010381"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(f"   ·   {roll}")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = MUTED

P("October 2026", size=10)
PB()

# ================================================================== TOC
H1("Contents")
add_toc()
P("If the page numbers are blank, right-click the table of contents and choose Update Field (Word) or Update (Pages).", size=8.5)
PB()

# ================================================================== 1
H1("1. Read this first")
CALLOUT(
    "The project in one paragraph",
    "Cogito is a web app that lets you upload documents and ask questions about them. It stores the documents and "
    "their searchable passages in MongoDB. When you upload a file, a background worker cuts it into passages, turns "
    "each passage into a list of numbers (an embedding), and saves those passages. When you ask a question, Cogito "
    "searches both by exact words and by meaning, combines the results, and asks a language model to answer using "
    "only the retrieved passages. Every answer lists the passages it used. The result is a searchable knowledge base "
    "that says where its answers came from.",
)

H2("1.1 The two-minute explanation")
P("Imagine a shared folder full of contracts, HR policies, privacy notices, and project notes. People know the answer is in there somewhere, but they cannot find it quickly. Ctrl+F only works if they already know the exact words. A chat assistant that reads the whole folder might sound confident and still be wrong, because it has no obligation to point at a source.")
P("Cogito fixes that in three steps. First, it indexes the documents automatically as soon as they are added. Second, it searches in two ways at once: exact words (good for clause numbers and names) and meaning (good for paraphrases). Third, it gives a short answer and lists the exact passages behind it. If the documents do not contain the answer, it should say so instead of inventing one.")
P("The project is built around MongoDB. MongoDB stores the documents, their passages, and the vectors used for semantic search. MongoDB also provides the keyword search, the vector search, the change stream that triggers indexing, the transactions that keep updates safe, and the aggregation pipelines behind the dashboard. Redis is a small, fast helper for rate limiting and for remembering where the worker left off. Next.js is the browser interface. FastAPI is the Python backend. A configurable language-model provider writes the final answer.")

H2("1.2 Five things to remember")
B([
    "MongoDB is the source of truth. Everything important — documents, chunks, vectors, logs, memories — lives there.",
    "A chunk is the unit of retrieval: one passage, its vector, and its metadata in one document.",
    "A change stream is the automatic trigger. The API writes a pending document; the worker reacts and indexes it.",
    "Hybrid search means keyword + meaning, fused into one ranking, then reranked.",
    "The answer is grounded: it comes with citations, and it should refuse when the sources do not support it.",
])

H2("1.3 A three-hour study plan")
TABLE(
    ["Time", "What to read", "What you should be able to do"],
    [
        ["0:00–0:40", "Part I (MongoDB, Redis, Next.js, FastAPI, RAG)", "Explain what each technology is and why it is in the project"],
        ["0:40–1:30", "Part II (architecture, data model, flows)", "Walk through an upload, a search, and an answer on a whiteboard"],
        ["1:30–2:10", "Part III (run, demo, troubleshooting)", "Start the stack and demo document CRUD + search + citations"],
        ["2:10–2:40", "Viva questions and glossary", "Answer the 20 questions without notes"],
        ["2:40–3:00", "Explain it in two minutes (section 4.4)", "Give a confident, simple overview"],
    ],
    [1.0, 2.4, 3.1],
)

H2("1.4 What you do not need to know")
P("You do not need to be a MongoDB administrator, a Redis expert, or a Next.js developer to understand this project. You need to know what each part is responsible for and how data moves between them. The guide gives you the mental model first, then the file names and code details.")
P("If you are asked a question you cannot answer, go back to the flow: upload → pending → change stream → worker → chunks → index → search → retrieve → answer → cite. Most questions can be answered by locating where you are in that flow.")

PB()

# ================================================================== PART I
H1("Part I · The technologies, from zero")

H2("2.1 What is NoSQL?")
P("A relational database (PostgreSQL, MySQL) stores data in tables with a fixed schema. You define columns first, and relationships are expressed with foreign keys and joins. That is excellent when the data is highly structured and reports need strict consistency across many tables. It becomes awkward when each record has a different shape, when the important read is a single aggregate, or when the data is a document with nested parts.")
P("NoSQL is an umbrella for databases that relax one or more of those assumptions. Document databases store JSON-like records; key-value stores map keys to values; graph databases store nodes and edges; wide-column stores organise data by rows with flexible columns. MongoDB is a document database. Redis is a key-value store. Cogito uses both, for different jobs.")
TABLE(
    ["Question", "Relational", "MongoDB (document)"],
    [
        ["Basic unit", "Row in a table", "Document in a collection"],
        ["Schema", "Fixed columns", "Flexible fields; shape can vary by document"],
        ["Relationships", "Foreign keys + joins", "Embedded documents or references"],
        ["Read pattern", "Join several tables", "Read one aggregate document"],
        ["Best fit", "Strict multi-table reporting", "Variable records + retrieval-style reads"],
    ],
    [1.4, 2.2, 2.9],
)

H2("2.2 MongoDB in 10 minutes")
P("MongoDB organises data in a simple hierarchy: a server contains databases; a database contains collections; a collection contains documents; a document contains fields. A document is a JSON-like object stored internally as BSON, which supports extra types such as dates and binary data. Every document has an _id field that uniquely identifies it.")
TABLE(
    ["Term", "Meaning", "In Cogito"],
    [
        ["Database", "A group of collections", "cogito"],
        ["Collection", "A group of documents, like a table", "documents, chunks, memories, query_logs"],
        ["Document", "One JSON-like record", "One source file or one passage"],
        ["Field", "A key-value pair inside a document", "title, text, embedding, status"],
        ["BSON", "MongoDB's binary JSON format", "Stores raw file bytes as Binary"],
        ["_id", "The unique identifier of a document", "A UUID string in this project"],
        ["Index", "A data structure that speeds up a query", "status, document_id, text, vector"],
    ],
    [1.1, 2.4, 3.0],
)
CODE(
    """{
  "_id": "82f9cb6d-...",
  "path": "contract_penalties.md",
  "title": "contract_penalties.md",
  "type": "contract",
  "status": "indexed",
  "chunk_count": 2,
  "summary": "Vendor penalties, force majeure, dispute resolution.",
  "created_at": 1792312802.17
}""",
    "A document in the documents collection. Optional fields simply appear when they are relevant.",
)
P("Flexible schema does not mean no design. It means optional fields do not require nullable columns or extra tables. A chunk may have a page number only when it came from a PDF; a document may have a summary only after enrichment. The application decides which fields are required, and MongoDB stores the rest without a migration.")

H3("Embedding versus referencing")
P("You can put related data inside a document (embedding) or store a reference to another document (referencing). Embed when the data is read together and belongs to one aggregate. Reference when the data is large, shared, or updated independently. Cogito embeds the vector, page, order, token count, and title inside each chunk, because retrieval reads all of them together. It references the parent document with document_id, because a document can have many chunks and the parent has its own lifecycle.")
CODE(
    """// Embedded: everything retrieval needs is in one document.
{ "_id": "chunk-1", "document_id": "doc-1", "text": "...",
  "embedding": [0.01, -0.04, ...], "page": 1, "title": "contract.md" }

// Referenced: the chunk points at its parent instead of copying the whole file.
{ "_id": "chunk-1", "document_id": "doc-1" }""",
    "The same idea in two shapes. Cogito uses the embedded version for chunks.",
)

H3("Indexes are a book index")
P("Without an index, MongoDB must inspect every document to answer a query. An index is a sorted structure that lets it jump to the matching documents. If you search for status = pending, an index on status turns a full scan into a lookup. If you search text, a text or search index tokenises the words. If you search vectors, an HNSW index finds nearby vectors without comparing every vector in the collection.")
P("Indexes are not free: each one uses disk and slows writes. That is why Cogito creates only the indexes its access patterns need: status, document_id, text, the vector index, (path, hash) for duplicate detection, (user_id, updated_at) for memory, and query_logs.ts for analytics.")

H3("Aggregation pipelines are an assembly line")
P("An aggregation pipeline is a list of stages. Each stage transforms the documents and passes them to the next stage, like stations on an assembly line. $match filters, $group aggregates, $sort orders, $limit cuts, $facet runs several sub-pipelines in one pass, and $bucketAuto groups numeric values into adaptive buckets. The dashboard is built from these stages, so the work happens inside MongoDB rather than in a Python loop.")

H3("Replica sets, change streams, and transactions")
P("A replica set is a group of MongoDB servers that keep copies of the same data. The primary accepts writes and records them in an operation log (oplog); secondaries replay the log. A change stream is a cursor over that log. It emits events when documents change. Cogito uses a change stream on the documents collection to trigger indexing. It also uses transactions, which let several writes succeed or fail as one unit. Both features require a replica set; the Docker setup creates a single-node replica set called rs0 so they work on one machine.")
CALLOUT(
    "What Cogito actually uses",
    "Document collections; flexible optional fields; embedded vectors; text and vector indexes; aggregation pipelines for analytics; a change stream for ingestion; multi-document transactions for re-index and delete; and a single-node replica set in Docker.",
    color="0F9D6E",
    fill=LIGHT_GREEN,
)

H2("2.3 Redis in 10 minutes")
P("Redis is an in-memory key-value store. You give it a key and a value, and it keeps them in RAM, which makes reads and writes extremely fast. It supports strings, lists, sets, sorted sets, hashes, and atomic operations. Many keys have a TTL (time to live): they disappear automatically after a set number of seconds. Redis is usually a helper, not the source of truth. If it restarts without persistence, the data can be lost; the main database must still be correct.")
P("Think of MongoDB as the filing cabinet and Redis as a stack of sticky notes on the desk. The cabinet is the durable record. The sticky notes are for things that are only useful for a short time: a rate-limit counter, a cached answer, or a bookmark that says where the worker stopped reading.")
P("Cogito uses Redis for three jobs.")
B([
    "Rate limiting. The key looks like rl:search:<first 16 characters of the query>. A sorted-set sliding window counts requests in the last 60 seconds. Above 60, the API returns HTTP 429 with a Retry-After header.",
    "Caching. The service writes short-lived keys such as answer:<id> with a TTL. cache_get and cache_set hide the Redis details and silently no-op when Redis is unavailable.",
    "Worker resume token. After each change-stream event, the worker stores the stream's resume token in Redis for seven days. On restart it resumes from that point instead of losing events that arrived while it was down.",
])
CODE(
    """# Local Redis path (simplified)
now = time.time()
pipe = redis.pipeline()
pipe.zremrangebyscore(key, 0, now - window)   # drop old hits
pipe.zadd(key, {member: now})                 # add this hit
pipe.zcard(key)                               # count the window
pipe.expire(key, window)
removed, added, count, _ = await pipe.execute()
return count <= limit, max(0, limit - count)""",
    "The rate limiter uses a sorted set. The real function is in app/telemetry.py.",
)
P("If Redis is down, Cogito still works. Rate limiting falls back to allow-all, caching becomes a no-op, and the worker logs a debug message. On the next start, the worker backfills any documents still marked pending, so no work is permanently lost. This graceful degradation is deliberate: Redis improves the experience but is not allowed to break the core application.")

H2("2.4 Next.js in 10 minutes")
P("React is a JavaScript library for building user interfaces from components. A component is a function that returns JSX, a syntax that looks like HTML inside JavaScript. Components hold state with hooks such as useState and run side effects with useEffect. When state changes, React re-renders the affected parts of the page.")
P("Next.js is a framework around React. It adds file-based routing, server-side rendering, optimised builds, and server route handlers. In the App Router (the style used here), a folder under app/ becomes a route. A file named page.tsx is the page for that route. A file named layout.tsx wraps the pages. A file named route.ts under app/api/ is a server-side API endpoint, not a page.")
TABLE(
    ["Next.js idea", "What it means", "In Cogito"],
    [
        ["Route", "A URL backed by a folder", "app/documents/page.tsx → /documents"],
        ["Layout", "Shared wrapper for pages", "app/layout.tsx"],
        ["Client component", "Runs in the browser; can use state", "documents/page.tsx, chat/page.tsx"],
        ["Server component", "Runs on the server by default", "home page, layouts"],
        ["Route handler", "Server-side API endpoint", "app/api/documents/route.ts"],
        ["Build", "Produces an optimised .next folder", "npm run build"],
    ],
    [1.35, 2.4, 2.75],
)
P("The browser never talks to FastAPI with the admin token. The documents page calls its own Next.js route handlers (/api/documents and /api/search). Those handlers run on the server, read COGITO_ADMIN_TOKEN, and forward the request to FastAPI with the X-Admin-Token header. This is why the token is not visible in the browser and why there is no CORS problem for the document CRUD calls.")
CODE(
    """// web/app/api/documents/route.ts (simplified)
export async function POST(req: Request) {
  const form = await req.formData();
  const res = await adminFetch("/ingest", { method: "POST", body: form });
  return NextResponse.json(await readJson(res), { status: res.status });
}""",
    "A Next.js route handler receives the browser request and forwards it server-side.",
)
P("The main pages are /documents (upload, list, edit, delete, search), /chat (ask questions and see citations), and /dashboard (aggregation panels). The documents page polls GET /api/documents every few seconds so a row moves from pending to indexed without a manual refresh. The chat page sends a question to FastAPI and renders the answer plus its citations. The dashboard fetches the analytics endpoints through a small proxy.")
P("To run the UI: npm install, then npm run dev for development on port 3000, or npm run build followed by npm start for a production build. NEXT_PUBLIC_ variables are exposed to the browser; server-only variables are not. That distinction is why the token lives in the server-side proxy.")

H2("2.5 FastAPI in 10 minutes")
P("FastAPI is a Python web framework. You decorate a Python function with @router.get or @router.post, and it becomes an HTTP endpoint. FastAPI reads the type hints, validates the request, and generates OpenAPI documentation at /docs. Pydantic models describe the request body; invalid input is rejected before your code runs. Async functions let the server wait on MongoDB or a model provider without blocking other requests.")
CODE(
    """class AskRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    k: int = Field(default=5, ge=1, le=20)

@router.post("/ask")
async def ask_question(body: AskRequest):
    return await ask(body.question, k=body.k)""",
    "A validated endpoint. Invalid questions never reach the retrieval code.",
)
P("The backend is split into routers (ingest, search, chat, memory, analytics, health) and core modules (retrieval, ingestion, LLM, evaluator, agent, memory, feedback). A request usually goes: router → service/core → MongoDB or model provider → JSON response. The ingest worker is a separate process that runs the same core code but is triggered by the change stream instead of an HTTP request.")

H2("2.6 LLMs, embeddings, and RAG")
P("A large language model (LLM) predicts text. It is very good at writing fluent answers, but it does not automatically know which document is true, and it can invent details. That failure is called hallucination. Cogito reduces it by grounding the model in retrieved passages and requiring citations.")
TABLE(
    ["Idea", "Plain-English meaning", "Where it appears"],
    [
        ["Embedding", "A list of numbers that represents meaning", "One vector per chunk and per question"],
        ["Vector", "The actual list, e.g. 768 floats", "Stored in chunks.embedding"],
        ["Cosine similarity", "How close two vectors point in the same direction", "Semantic search fallback"],
        ["Chunk", "A passage small enough to retrieve and cite", "chunks collection"],
        ["Retrieval", "Finding the most relevant chunks", "Hybrid search"],
        ["RAG", "Retrieve first, then generate the answer", "The /ask pipeline"],
        ["Citation", "The passage behind a claim", "citations array in the answer"],
        ["Judge", "A second model or heuristic that scores the answer", "eval_results"],
    ],
    [1.2, 2.7, 2.6],
)
P("RAG stands for retrieval-augmented generation. Think of an open-book exam: the model may consult the retrieved pages, but it must answer the question and show which pages it used. Cogito's prompt tells the model to answer concisely, cite only the chunks it used, and return a refusal when the context does not contain the answer.")
P("Cogito is not tied to one model vendor. The configuration selects an LLM provider and an embedding provider; the code supports OpenAI-compatible providers, Ollama, and a built-in local fallback. The local fallback keeps the pipeline runnable offline. The model is replaceable; the database contract is not. That is an important architecture point: the LLM writes the answer, but MongoDB decides what evidence is available.")

PB()

# ================================================================== PART II
H1("Part II · How Cogito works")

H2("3.1 The product")
P("Cogito is a document assistant. A user uploads a PDF, DOCX, Markdown file, HTML page, CSV, JSON file, image, or plain text. The app stores it, indexes it, and makes it searchable. The user can then search for passages or ask a question in natural language. The answer is short, grounded in the retrieved passages, and accompanied by citations. The user can also edit the document metadata or content, re-index it, delete it, and inspect the chunks that were created.")
P("The knowledge base starts empty. That is deliberate: during the demo, every document is added live, so the evaluator can see the pending → indexed transition and the search results change. There is no hidden sample content.")
P("The real-world use case is any document-heavy team: legal and compliance teams looking up clauses, HR teams answering policy questions, operations teams searching notices, or auditors checking evidence. The same design works for a personal knowledge base or an internal help centre.")

H2("3.2 Architecture")
FIG("architecture.png", "Figure 1. The runtime architecture. The browser talks to Next.js; Next.js talks to FastAPI; FastAPI and the worker talk to MongoDB, Redis, Postgres, and the model provider.")
P("The Docker Compose stack runs MongoDB (with mongot, the search component), the FastAPI backend, the ingest worker, Redis, and Postgres. The Next.js UI runs separately with npm on port 3000. In a deployment you could containerise the UI too, but keeping it separate makes frontend development simpler.")
TABLE(
    ["Service", "Port", "Responsibility", "Durable data?"],
    [
        ["Next.js UI", "3000", "Pages, forms, chat, dashboard, server-side API proxies", "No"],
        ["FastAPI backend", "8000", "REST API, validation, retrieval, answer generation", "No"],
        ["Ingest worker", "—", "Watches the change stream; chunks and embeds documents", "No"],
        ["MongoDB + mongot", "27018 on host; 27017 in Docker", "Documents, chunks, vectors, logs; search and aggregation", "Yes"],
        ["Redis", "6379", "Rate limiting, short-lived cache, worker resume token", "No (helper)"],
        ["Postgres", "5432", "Optional analytics mirror", "Optional"],
    ],
    [1.4, 1.5, 3.1, 1.0],
)
P("The backend never reads files from the host's filesystem. The uploaded bytes are stored in the documents collection as BSON Binary, so the worker can process the document from the change event alone. This is why a document can be indexed even if the API and worker are separate processes or containers.")
CALLOUT(
    "The most important line in the architecture",
    "The API only writes a pending document. Everything else — chunking, embedding, indexing, status updates — is triggered by MongoDB's change stream. There is no scheduler and no manual re-run.",
    color="B45309",
    fill=LIGHT_AMBER,
)

H2("3.3 Data model")
FIG("data_model.png", "Figure 2. The core collections. documents and chunks are the retrieval path; the other collections support chat history, memory, logging, and evaluation.")
P("Cogito uses seven collections. The first two are the heart of the application.")
TABLE(
    ["Collection", "Stores", "Key fields"],
    [
        ["documents", "Source files and metadata", "_id, path, title, type, hash, status, raw, summary, keywords, chunk_count"],
        ["chunks", "Retrievable passages and vectors", "_id, document_id, text, embedding, page, order, token_count, title, boost"],
        ["conversations", "Chat history", "_id, user_id, messages[], created_at"],
        ["memories", "Long-term facts", "_id, user_id, fact, embedding, importance, updated_at"],
        ["query_logs", "One row per question", "question, chunks_returned, chunk_ids[], latency_ms, answer_id, ts"],
        ["eval_results", "Judge scores", "answer_id, faithfulness, relevance, judge, ts"],
        ["feedback", "Ratings (API only)", "answer_id, rating, comment, ts"],
    ],
    [1.15, 2.15, 3.4],
)
P("The key modeling decision is the chunk. A relational schema might store the passage in one table, the vector in another, and the document metadata in a third; retrieval would join them. Here the chunk contains the passage, its vector, its page, its order, its token count, and a copy of the document title. One read gives the retrieval code everything it needs. The title is deliberately duplicated so a citation does not need a second lookup; renaming a document also updates its chunks.")
CODE(
    """{
  "_id": "c392cfe9-...",
  "document_id": "82f9cb6d-...",
  "text": "4.1 Late Delivery. The vendor agrees ...",
  "embedding": [0.013, -0.041, 0.008, ...],   // 768 numbers
  "page": 1,
  "order": 3,
  "token_count": 141,
  "title": "contract_penalties.md",
  "mime": "text/markdown",
  "boost": 1.0,
  "created_at": 1792312802.17
}""",
    "One chunk documents the whole retrieval contract: text, vector, page, title, and a ranking boost.",
)
P("The relationships are simple. A document has many chunks (document_id). A conversation embeds its messages in one array. A memory is a standalone fact with its own embedding. Query logs reference chunk IDs and an answer ID; evaluation results and feedback reference the answer ID. These are references, not enforced foreign keys, because MongoDB does not require a join across collections for the read path.")
H3("Indexes and access patterns")
TABLE(
    ["Access pattern", "Query", "Index"],
    [
        ["Find pending work", "{ status: \"pending\" }", "documents.status"],
        ["Fetch a document's chunks", "{ document_id }", "chunks.document_id"],
        ["Keyword search", "$search or $text", "MongoDB Search index / text index"],
        ["Semantic search", "$vectorSearch", "HNSW vector index, cosine, 768d"],
        ["Detect duplicate uploads", "{ path, hash }", "compound (path, hash)"],
        ["Recall a user's memories", "{ user_id }", "compound (user_id, updated_at)"],
        ["Analytics window", "{ ts }", "query_logs.ts"],
    ],
    [1.9, 2.1, 3.0],
)
P("Scalability is a design concern, not a claim that this demo is a production cluster. The natural shard candidate is document_id, because chunks for one document are read together and can be co-located. Read replicas can serve search and analytics. Vectors can use scalar quantization to reduce memory. The vector search over-requests candidates (100 for 5 results) so the fused ranking has enough material; that is a recall/latency trade-off, not a guarantee of perfect recall.")

H2("3.4 The upload and indexing flow")
FIG("upload_flow.png", "Figure 3. Uploading a document is a short API write followed by an asynchronous worker pipeline.")
N([
    "The user selects a file or pastes text on /documents. The browser sends it to the Next.js route handler POST /api/documents.",
    "The route handler forwards the multipart request to FastAPI POST /ingest with the admin token. FastAPI validates the size, type, and filename, then calls create_document.",
    "create_document inserts a document with status = pending and the raw bytes. The API returns the new document ID immediately; it does not wait for indexing.",
    "MongoDB's change stream emits an insert event. The ingest worker is watching the documents collection with full_document = updateLookup, so it receives the entire pending document.",
    "The worker extracts text according to the file type (PDF, DOCX, HTML, CSV, JSON, plain text, OCR for images), splits it into overlapping chunks, and sends the chunk texts to the embedding provider.",
    "The worker writes the chunks and flips the document to indexed inside one multi-document transaction. If the transaction is unavailable, it falls back to the same writes without a transaction and logs the reason.",
    "The UI polls GET /api/documents every few seconds. The row changes from pending to indexed and the chunk count appears.",
])
TABLE(
    ["Status", "Meaning", "What to do if stuck"],
    [
        ["pending", "Queued; the worker has not finished", "Check that the worker container is running"],
        ["indexed", "Chunks are stored and searchable", "Open the document and inspect its chunks"],
        ["failed", "Extraction or embedding failed", "Open the document to see the stored error"],
    ],
    [1.0, 2.7, 2.9],
)
P("Two reliability details matter. First, the worker saves the change stream's resume token in Redis after every event, so a restart continues from the last processed event. Second, on startup the worker backfills any documents still marked pending, which covers the case where the API accepted a document while the worker was down. The transaction ensures that a crash in the middle of re-indexing never leaves a document marked indexed with missing or half-written chunks.")
P("The main code locations are create_document and process_document in src/backend/app/core/ingestion.py, the watch loop in src/backend/app/ingest_worker.py, and _commit_document for the transaction. If you remember one file for this flow, remember ingestion.py.")

H2("3.5 The search flow")
FIG("search_flow.png", "Figure 4. A question is embedded once, matched lexically and semantically, fused, and reranked.")
N([
    "The question arrives at FastAPI POST /search (or inside POST /ask). The backend embeds the question with the configured embedding provider.",
    "Lexical search matches exact words and identifiers with $search (mongot) or $text (fallback). This is strong for clause numbers, names, and exact phrases.",
    "Semantic search uses $vectorSearch (mongot) or brute-force cosine (fallback). This is strong for paraphrases: “late delivery fees” can match “penalty for delay”.",
    "The two ranked lists are fused. In full mode MongoDB's $rankFusion does it; in fallback mode a Python RRF function does it. RRF scores each document by 1/(60 + rank) and sums the contributions. It uses ranks, not raw similarity scores, so the two branches do not need to be on the same scale.",
    "The fused candidates are reranked by a cross-encoder when the optional library is installed; otherwise a lexical overlap boost promotes chunks that share distinctive terms with the question.",
    "The top-k chunks are returned. The same fields are returned in both modes: _id, text, document_id, title, page, score.",
])
TABLE(
    ["Condition", "Path", "Operators"],
    [
        ["MongoDB Search (mongot) available", "Full hybrid", "$search + $vectorSearch + $rankFusion"],
        ["Plain mongod (no mongot)", "Fallback", "$text + cosine similarity + Python RRF"],
    ],
    [2.6, 1.4, 2.7],
)
P("The system probes for mongot at startup and caches the result. You can see which path is active at GET /capabilities. The fallback is not a stub: it returns the same result shape and powers the same UI. In the Docker stack, mongot is bundled with MongoDB, so the native pipeline is the default. If you point MONGODB_URI at a plain mongod, the fallback keeps the app usable.")
P("Query optimization lives in three places: the indexes (text, HNSW vector, compound), the candidate over-request (search_num_candidates = 100 for k = 5, roughly 20× the result limit), and the reranker. The 20× guidance is about recall: approximate vector search may miss a relevant neighbor if you ask for too few candidates before fusion.")

H2("3.6 The ask/answer flow")
P("The /ask endpoint wraps retrieval with generation. The sequence is: retrieve chunks → build a prompt containing the chunks and the question → call the configured LLM → parse a JSON object with an answer and citations → log the question and latency → evaluate the answer → return it to the UI.")
CODE(
    """{
  "answer": "Late delivery accrues a 0.5% penalty per day, capped at 15%.",
  "citations": [
    { "document_id": "82f9cb6d-...",
      "title": "contract_penalties.md",
      "excerpt": "4.1 Late Delivery. The vendor agrees ...",
      "score": 0.079 }
  ]
}""",
    "The answer contract. The UI renders the answer and one citation card per source.",
)
P("The system prompt is the guardrail. It instructs the model to answer only from the context, keep the answer to at most four sentences, cite only the chunks it used, keep excerpts short, and return exactly “I could not find this in the knowledge base.” with no citations when the context does not support an answer. This is why an unrelated question should be refused rather than answered from general knowledge.")
P("The evaluation step is a second, cheaper judgment. With a model configured, the judge scores faithfulness (are the claims supported by the retrieved chunks?) and relevance (does the answer address the question?). Without a model, a deterministic local heuristic does the same job. The scores are stored in eval_results and shown on the dashboard, so changes in retrieval or prompting are visible over time.")
P("The agent endpoint adds a small loop: retrieve, compute confidence as the mean score of the top three chunks, refine the query with distinctive terms from the best chunk if confidence is weak, retrieve again (bounded by AGENT_MAX_ITERATIONS), then generate. The multi-query endpoint expands the question into several rewrites, retrieves for each, and fuses the results. Both reuse the same retriever and answer contract; they are optional paths, not separate architectures.")
CALLOUT(
    "How hallucination is controlled",
    "Retrieval limits what the model can see. The prompt requires citations. The answer can refuse. The judge scores faithfulness. No single mechanism is perfect, but together they make unsupported answers visible and reviewable.",
    color="6D4AFF",
    fill="F3EDFF",
)

H2("3.7 Update, delete, and memory")
P("CRUD is not limited to upload and search. The documents page and the API expose all four operations.")
TABLE(
    ["Operation", "How", "What happens internally"],
    [
        ["Create", "POST /ingest or /ingest/raw", "Insert pending document; change stream indexes it"],
        ["Read", "GET /ingest/documents, /{id}, /chunks; POST /search", "Metadata, text, chunks, ranked passages"],
        ["Update title", "PATCH /ingest/documents/{id}", "Updates the document and the denormalized title on its chunks"],
        ["Update content", "PATCH /ingest/documents/{id} with text", "Replaces raw bytes/hash, marks pending; worker re-chunks and re-embeds"],
        ["Re-index", "POST /ingest/documents/{id}/reindex", "Marks pending; useful after changing the embedding provider"],
        ["Delete", "DELETE /ingest/documents/{id}", "Deletes the document and all chunks in one transaction"],
    ],
    [1.1, 2.6, 3.0],
)
P("The update path is a good example of the data model in action. Changing the title updates the document and every chunk's title, because citations read the title from the chunk. Changing the text replaces the stored raw bytes and the content hash, sets status back to pending, and clears any previous error. The worker then rebuilds the chunks; the old chunks are deleted inside the same transaction that inserts the new ones.")
P("Delete is transactional for the same reason: removing the document but leaving its chunks would create orphaned search results. The delete route deletes the chunks and the document in one session/transaction, with a fallback if transactions are unavailable. After deletion, GET /ingest/documents/{id} returns 404 and a search no longer returns those chunks.")
P("Memory is a separate feature for long-term facts. POST /memories stores a fact with its embedding; GET /memories lists them; POST /memories/recall computes cosine similarity against the user's memories and returns the closest ones. The chat endpoint can inject recalled facts into the prompt. PATCH and DELETE let a fact be corrected or removed. This is the “agent memory” pattern implemented on the same database, not a separate vector service.")

H2("3.8 Analytics")
P("Every question writes a row to query_logs with the question, the number of chunks returned, the chunk IDs, the latency, the answer ID, and a timestamp. The dashboard is built from that log and from eval_results and feedback.")
CODE(
    """{ $facet: {
    totals:     [{ $count: "n" }],
    unanswered: [{ $match: { chunks_returned: 0 } }, { $count: "n" }],
    latency:    [{ $group: { _id: null, avg: { $avg: "$latency_ms" } } }],
    by_day:     [{ $group: { _id: { $dateToString: {
                    format: "%Y-%m-%d",
                    date: { $toDate: { $multiply: ["$ts", 1000] } } } },
                    count: { $sum: 1 } } }]
}}""",
    "One $facet pass computes several dashboard metrics at once. Sub-pipelines return arrays; the API unwraps the first element.",
)
P("The other panels use $bucketAuto for an eight-bucket latency histogram and $group + $sort + $limit for the most frequent questions. The dashboard also shows document/chunk counts, ingestion freshness, and the judge's average faithfulness/relevance. The point is not the specific charts; it is that the database computes the analytics, so the application does not download raw collections and loop over them in Python.")

H2("3.9 Redis and Postgres in the running system")
P("Redis is a supporting service with three concrete jobs. The rate-limit key is rl:search:<query prefix>; a sorted set stores timestamps, and the API returns 429 when the 60-per-minute window is full. The cache helpers store short-lived values such as answer IDs. The worker stores cogito:resume_token so it can resume the change stream after a restart. Redis being unavailable degrades these features but does not stop ingestion, search, or answering.")
P("Postgres is an optional analytics mirror. The backend can write evaluation/feedback records to Supabase or the bundled Postgres through the integrations layer. It is not on the critical path: the MongoDB collections are the source of truth, and the app runs without Postgres. This is a useful distinction in the viva: not every service in the compose file is equally important.")
TABLE(
    ["Failure", "What still works", "What degrades"],
    [
        ["Redis down", "Ingestion, search, answers, CRUD", "Rate limit allows all; cache misses; resume token not persisted"],
        ["Postgres down", "Everything core", "Analytics mirror writes are skipped"],
        ["mongot absent", "Search and answers via fallback", "No native $search/$vectorSearch/$rankFusion"],
        ["Worker down", "Uploads and reads", "New documents stay pending until the worker returns"],
        ["Model provider rate-limited", "Retrieval and citations", "Answer generation may return a retry message"],
    ],
    [1.6, 2.5, 2.6],
)

H2("3.10 Frontend tour")
P("The Next.js app has three user-facing pages and a few server route handlers.")
B([
    "/documents: the main control panel. It lists documents, shows status and chunk counts, opens a detail view, and provides upload, paste, search, rename, content editing, re-index, and delete.",
    "/chat: asks a question and renders the answer with citation cards. It has single-pass, agent, and multi-query modes. The old thumbs-up/down reactions have been removed.",
    "/dashboard: renders the aggregation results with Recharts: totals, top questions, query volume by day, latency, eval scores, and freshness.",
    "app/api/documents/*: server-side handlers for list/create/read/update/delete/reindex/chunks. They add the admin token and forward to FastAPI.",
    "app/api/search/route.ts: a server-side proxy for the search endpoint.",
    "lib/backend.ts: the backend URL and a small fetch helper; lib/server.ts: the token-aware server fetch.",
])
P("The documents page is a client component because it uses useState, useEffect, and useRef. Its state machine is simple: load the list, poll while anything is pending, upload a file or pasted text, open a document to fetch its metadata and chunks, edit title/text, re-index, or delete. Every mutation calls the Next.js route handler, then reloads the list so the UI reflects MongoDB, not a local guess.")
P("To add a page, create a folder under web/app with a page.tsx. To add an API proxy, create a route.ts under web/app/api. To style, use the existing classes in web/app/globals.css. The layout and navigation are shared through components/nav.tsx. If a page shows stale data, check the browser Network tab and the Next.js server log; if a write fails, check the token in web/.env.local against ADMIN_TOKEN in the root .env.")

H2("3.11 Backend tour")
P("The backend lives in src/backend/app. The main.py file creates the FastAPI app, wires the routers, configures CORS, and starts the index/search probe on startup. The routers translate HTTP requests into calls to core modules.")
TABLE(
    ["Path", "What lives there"],
    [
        ["app/main.py", "FastAPI app, lifespan, CORS, routers, 429 handler"],
        ["app/config.py", "Environment settings: Mongo URI, providers, dimensions, limits"],
        ["app/db.py", "Motor client, collection handles, index creation, search-mode probe"],
        ["app/deps.py", "Admin-token dependency"],
        ["app/routers/ingest.py", "Upload, raw text, list, detail, chunks, update, reindex, delete"],
        ["app/routers/search.py", "Search, multi-query search, ask, agent, compare"],
        ["app/routers/chat.py", "Chat with history and memory recall"],
        ["app/routers/memory.py", "Memory CRUD and feedback"],
        ["app/routers/analytics.py", "Aggregation endpoints for the dashboard"],
        ["app/routers/health.py", "Health, capabilities, webhooks"],
        ["app/core/ingestion.py", "Create/process/commit documents; the transaction"],
        ["app/core/hybrid_search.py", "Full and fallback retrieval, RRF, reranking hook"],
        ["app/core/llm.py", "Provider clients, prompt, cited-answer contract"],
        ["app/core/evaluator.py", "LLM judge and local heuristic judge"],
        ["app/core/agent.py", "LangGraph retrieve → refine → generate loop"],
        ["app/core/memory.py", "Remember/recall/update/delete facts"],
        ["app/ingest_worker.py", "Change-stream watch loop and backfill"],
        ["app/scripts/ensure_indexes.py", "One-shot index provisioning"],
    ],
    [2.0, 4.5],
)
P("A typical request is easy to trace. POST /ask enters routers/search.py, calls core/service.py:ask, which calls hybrid_search, then generate_cited_answer, then log_query and evaluate. POST /ingest enters routers/ingest.py, calls core/ingestion.py:create_document, and returns. The worker later calls process_document and _commit_document. If you can name those functions, you can answer most “where does this happen?” questions.")

H2("3.12 Search modes and capabilities")
P("GET /capabilities is the honest status endpoint. It reports whether mongot is available (full_hybrid), which model/embedding providers are configured, and which optional integrations are enabled. The UI uses it for the home-page badges. During the viva, call it and explain the difference between the native and fallback paths.")
P("In full mode, the retrieval aggregation runs $search and $vectorSearch and fuses them with $rankFusion. In fallback mode, the lexical branch uses a text index and the semantic branch computes cosine similarity in Python, then a small RRF function fuses the lists. The result contract is identical, which is why the rest of the application does not care which mode is active. The fallback is slower on large collections, but it keeps the demo portable.")

H2("3.13 Performance")
TABLE(
    ["Measurement", "Result", "Why"],
    [
        ["Ingestion, local embedder, first run", "0.26 s", "Small documents; local hashing embedder"],
        ["Ingestion, warm local", "6 ms", "No network call"],
        ["Hosted embedding", "1.6–4.4 s", "Network round-trip per document"],
        ["Answer with hosted model + judge", "about 7.4 s", "Two model calls dominate"],
        ["Automated tests", "17 passing", "Unit tests; no network"],
        ["Vector candidates", "100 for k=5", "Recall margin for fusion"],
    ],
    [2.5, 1.2, 3.0],
)
P("These numbers are from the verification run, not a production benchmark. The database work is fast; hosted model calls are the slow part. That is why the ingestion worker exists: it keeps model latency out of the upload request. The answer path still calls the model synchronously, so latency there is expected. The honest statement is “the architecture isolates the slow parts; it does not make network calls free.”")

H2("3.14 Security, validation, and failure behavior")
P("The write path is protected by a shared admin token. The browser never receives it: the Next.js route handlers run on the server, read COGITO_ADMIN_TOKEN, and add X-Admin-Token when forwarding to FastAPI. FastAPI compares the token in constant time. Read endpoints are open so the demo works without accounts; that is a deliberate limitation, not an oversight.")
B([
    "Pydantic validates request bodies: question lengths, k bounds, memory facts, and update fields.",
    "Uploads are limited to 10 MB and a whitelist of document types; empty files are rejected.",
    "The rate limiter returns 429 with Retry-After when the per-query window is full.",
    "Raw file bytes are never returned by the document list/detail endpoints.",
    "CORS is configured through the CORS_ORIGINS setting; the default is permissive for local development.",
    "Errors are mapped to readable messages in the UI: 400, 401, 404, 413, 422, 429.",
])
P("The limitations are equally important to state: there is no per-user login or document ACL yet, the demo runs over HTTP, the database is a single node, and the local rate limiter depends on Redis. These are the first things to improve in a production version: OAuth/JWT, per-user access control, TLS, a managed MongoDB deployment with backups, and an automated database integration suite in CI.")

PB()

# ================================================================== PART III
H1("Part III · Run it, demo it, defend it")

H2("4.1 Run the project")
P("Prerequisites: Docker Desktop (with enough disk space), Node.js 20+, and this repository. The first run builds the backend images and pulls MongoDB, Redis, and Postgres; it can take a few minutes.")
CODE(
    """# 1. Configure the environment (the file is gitignored; never commit it)
cp .env.example .env
# edit .env: choose LLM_PROVIDER / EMBEDDING_PROVIDER and put keys there if needed

# 2. Start the backend stack
docker compose up -d --build

# 3. Check that the services are healthy
docker compose ps
curl -s localhost:8000/health
curl -s localhost:8000/capabilities

# 4. Start the web UI (separate terminal)
cd web
npm install
npm run dev
# open http://localhost:3000/documents""",
    "The normal development workflow.",
)
P("Ports: the UI is 3000, FastAPI is 8000, MongoDB is 27018 on the host (27017 inside Docker), Redis is 6379, and Postgres is 5432. MongoDB is published on 27018 to avoid clashing with a local mongod. The web UI is not part of the compose file; it runs with npm. If you prefer a production build, use npm run build then npm start.")
P("Start with an empty knowledge base. Click Browse files (or drag a file) on /documents, choose a type, and click Ingest document. The row appears as pending. Watch it become indexed, then open it to see the extracted text and chunks. Search for a phrase from the document; then ask a question in /chat. If the answer is not in the documents, ask something unrelated and confirm that the app refuses with no citations.")
TABLE(
    ["Symptom", "Likely cause", "Fix"],
    [
        ["Upload returns 401", "Web token does not match backend", "Set COGITO_ADMIN_TOKEN in web/.env.local to ADMIN_TOKEN"],
        ["Document stays pending", "Worker is not running", "docker compose ps; start ingest-worker"],
        ["No search results", "Document not indexed yet", "Open the document and check chunks/status"],
        ["hybrid_search false", "mongot unavailable", "Use the bundled atlas-local image; fallback still works"],
        ["429 from /ask", "Model or search rate limit", "Wait for Retry-After; free tiers are limited"],
        ["Web page cannot reach API", "Wrong NEXT_PUBLIC_COGITO_API", "Point it at http://localhost:8000"],
        ["File dialog does not open", "Embedded preview blocks native dialogs", "Open localhost:3000 in Chrome/Safari, or use Paste text"],
    ],
    [1.7, 2.2, 2.8],
)
P("To stop the stack, run docker compose down. To reset the knowledge base without deleting the database volume, use the delete operations in the UI or a mongosh deleteMany on documents and chunks. To stop everything including the web UI, close the npm process.")

H2("4.2 The eight-minute demo")
P("The demo starts from an empty knowledge base. Do not pre-load sample documents; the point is to show the state transition live.")
N([
    "Show /documents with 0 documents and 0 chunks. Say: “The knowledge base is empty; everything you see next is created live.”",
    "Upload one PDF/DOCX/Markdown file. Point at status pending. Say: “The API only wrote a pending document; MongoDB's change stream wakes the worker.”",
    "Wait for indexed. Open the document and show the chunks. Say: “The worker extracted, chunked, and embedded the text, then committed the chunks and status in one transaction.”",
    "Search for a literal phrase, then a paraphrase. Say: “Keyword and vector search run together and are fused; the scores come from MongoDB.”",
    "Ask a question in /chat. Show the answer and citations. Ask an unrelated question and show the refusal with no citations. Say: “This is the grounding guard.”",
    "Rename the document, edit its content, re-index it, then delete it. Show the dashboard. Say: “This is full CRUD plus aggregation analytics.”",
])
P("If a service is slow, narrate what should happen instead of waiting silently. If the model is rate-limited, show the retry message and explain the backoff; the retrieval and citations still work. If the demo is offline, the local fallback still produces the same response shape.")

H2("4.3 Viva questions")
P("These are short answers. The full, more conversational versions are in docs/VIVA_PREP.md.")
TABLE(
    ["Question", "Short answer"],
    [
        ["What is the project?", "A document assistant: upload files, search them, ask questions, and get cited answers."],
        ["Why MongoDB?", "A chunk is one document with text, vector, and metadata; MongoDB also supplies search, vectors, change streams, transactions, and aggregation."],
        ["Why not Postgres + pgvector?", "It is viable, but retrieval would need joins plus a separate lexical index and more moving parts."],
        ["What is a change stream?", "A cursor over the replica set's oplog that emits events when documents change."],
        ["Why a replica set?", "Change streams and transactions require the oplog, which exists with replication."],
        ["What is a chunk?", "A passage plus its vector, page, order, token count, and denormalized title."],
        ["Why store the vector in the chunk?", "Retrieval reads text, vector, and metadata together; no join is needed."],
        ["What is hybrid search?", "Keyword search plus semantic vector search, fused into one ranking."],
        ["What is RRF?", "Reciprocal Rank Fusion: sum 1/(60 + rank) across ranked lists; ranks avoid comparing raw scores."],
        ["Why 100 candidates for 5 results?", "Approximate search can miss neighbors; over-requesting protects recall before fusion."],
        ["What does Redis do?", "Rate limiting, a short-lived cache, and the worker's change-stream resume token."],
        ["What if Redis is down?", "Rate limiting allows all, cache misses, and the worker relies on startup backfill."],
        ["What is Next.js?", "A React framework with file-based routes, server components, and server route handlers."],
        ["Why a Next.js proxy?", "It injects the admin token server-side, so the browser never sees it and CORS is avoided."],
        ["How is auth handled?", "A shared admin token on writes, constant-time compare, server-side proxy; no per-user login yet."],
        ["How is input validated?", "Pydantic lengths/bounds, a 10 MB upload cap, a type whitelist, and empty-value checks."],
        ["How do you test it?", "17 unit tests, Ruff + build in CI, and a documented upload-to-delete E2E run."],
        ["How do you prevent hallucination?", "Retrieve only relevant chunks, require citations, allow a refusal, and score faithfulness."],
        ["What is the fallback search mode?", "$text + cosine similarity + Python RRF when mongot is absent."],
        ["What would you improve?", "Per-user auth/ACL, a Mongo integration suite in CI, TLS/HA, and a production deployment."],
    ],
    [2.5, 4.3],
    font_size=8.4,
)

H2("4.4 Explain it in two minutes")
P("Use this script when someone asks “what did you build?”. Do not read it; practice it aloud.")
P("“Cogito is a document assistant built on MongoDB. You upload a PDF or a policy document, and it indexes itself: an API writes a pending record, a MongoDB change stream wakes a worker, and the worker splits the text into chunks, turns each chunk into a vector, and stores the chunks and the status update in one transaction. When you ask a question, Cogito searches in two ways: keyword search for exact terms and clause numbers, and vector search for meaning. It fuses the two rankings, reranks the candidates, and gives the top passages to a language model with a strict instruction: answer only from these passages and cite them. If the passages do not contain the answer, it refuses. The knowledge base, the search, the vectors, and the analytics all live in MongoDB; Redis is only a fast helper for rate limiting and resume tokens. The project also has full document CRUD and an aggregation dashboard.”")
P("Then invite one question. The most likely follow-up is “why MongoDB?”. Answer with the chunk aggregate and the fact that one deployment provides search, vectors, events, transactions, and analytics.")

H2("4.5 Glossary")
TABLE(
    ["Term", "Definition"],
    [
        ["NoSQL", "A family of databases that do not require fixed relational tables."],
        ["Document", "A JSON-like record; the unit of storage in MongoDB."],
        ["Collection", "A group of documents; roughly a table."],
        ["BSON", "MongoDB's binary encoding of JSON plus extra types."],
        ["Index", "A structure that speeds up matching queries."],
        ["Aggregation", "A pipeline of stages that transforms documents."],
        ["Replica set", "A group of MongoDB nodes that replicate the oplog."],
        ["Change stream", "A cursor that emits database change events."],
        ["Transaction", "Several writes that commit or abort together."],
        ["Embedding", "A numeric vector representing meaning."],
        ["Vector search", "Finding nearby vectors by similarity."],
        ["Cosine similarity", "A measure of the angle between two vectors."],
        ["Chunk", "A passage used as the retrieval unit."],
        ["RAG", "Retrieval-augmented generation: retrieve evidence, then answer."],
        ["LLM", "A large language model that generates text."],
        ["Prompt", "The instructions and context given to the model."],
        ["Citation", "The source passage behind an answer."],
        ["Hallucination", "A confident but unsupported model claim."],
        ["TTL", "A key's time to live before Redis deletes it."],
        ["Rate limit", "A cap on requests in a time window."],
        ["React", "A component-based JavaScript UI library."],
        ["Next.js", "A React framework with routing and server handlers."],
        ["FastAPI", "A Python web framework with validation and OpenAPI."],
        ["Pydantic", "Python data validation used by FastAPI."],
        ["Async", "Concurrent code that waits without blocking a thread."],
    ],
    [1.4, 5.2],
    font_size=8.6,
)

H2("4.6 File map")
CODE(
    """Cogito/
├── docker-compose.yml              MongoDB, backend, worker, Redis, Postgres
├── src/backend/
│   ├── app/
│   │   ├── main.py                 FastAPI app + routers
│   │   ├── config.py               settings and provider selection
│   │   ├── db.py                   Mongo client, indexes, search probe
│   │   ├── deps.py                 admin-token dependency
│   │   ├── ingest_worker.py        change-stream worker + backfill
│   │   ├── routers/                ingest, search, chat, memory, analytics, health
│   │   └── core/                   ingestion, hybrid_search, llm, evaluator, agent, memory
│   └── tests/                      unit tests
├── web/                            Next.js UI
│   ├── app/documents/              upload, list, edit, delete, search
│   ├── app/chat/                   cited answers
│   ├── app/dashboard/              aggregation panels
│   ├── app/api/                    server-side proxies
│   ├── components/                 shared navigation
│   └── lib/                        backend URL + token-aware fetch
├── docs/                           report, feature map, viva, demo, deck, this guide
└── docker/                         MongoDB bootstrap + Postgres schema""",
    "A trimmed tree. The README has the full layout.",
    size=8.0,
)

H2("4.7 Submission checklist")
B([
    "Course report (DOCX/PDF): problem, design, implementation, results, limitations.",
    "Feature map: docs/DATABASE.md — every MongoDB feature mapped to code.",
    "Viva sheet: docs/VIVA_PREP.md.",
    "Demo script: docs/DEMO_SCRIPT.md.",
    "Test evidence: docs/TEST_REPORT.md.",
    "Presentation: docs/Cogito_Presentation.pptx (11 slides).",
    "This guide: docs/Cogito_Project_Guide.docx.",
    "Plagiarism report: required by the rubric and generated by the institution (Turnitin/DrillBit). The repository cannot create it.",
])
CALLOUT(
    "Before you walk in",
    "Run the stack once from a clean start, upload one document, and complete the full demo path. A working demo is more convincing than any slide. If the model provider is rate-limited, the refusal and citations still demonstrate the pipeline.",
    color="B45309",
    fill=LIGHT_AMBER,
)

PB()

# ================================================================== appendix
H1("Appendix · Commands and code references")
H2("A.1 Commands cheat sheet")
CODE(
    """# Stack
docker compose up -d --build
docker compose ps
docker compose logs -f backend
docker compose logs -f ingest-worker
docker compose down

# Health and capability
curl -s localhost:8000/health
curl -s localhost:8000/capabilities

# A first document
curl -s -X POST localhost:8000/ingest/raw \\
  -H 'Content-Type: application/json' -H 'X-Admin-Token: change-me' \\
  -d '{"text":"Annual leave is 20 days.","filename":"policy.txt","type":"text"}'

# Ask
curl -s -X POST localhost:8000/ask \\
  -H 'Content-Type: application/json' \\
  -d '{"question":"How much annual leave is there?","k":5}'

# Tests
cd src/backend && pytest tests/ -q
cd src/backend && ruff check app
cd web && npm run build""",
    "Use the token from your own .env, not the placeholder.",
    size=8.2,
)

H2("A.2 Where to look in the code")
TABLE(
    ["If you want to understand…", "Open…"],
    [
        ["How a document is created", "app/core/ingestion.py · create_document"],
        ["How the worker reacts", "app/ingest_worker.py · watch"],
        ["How chunks are written atomically", "app/core/ingestion.py · _commit_document"],
        ["How hybrid search is built", "app/core/hybrid_search.py · _full_hybrid / _fallback_hybrid"],
        ["How the answer prompt works", "app/core/llm.py · SYSTEM_PROMPT / generate_cited_answer"],
        ["How the judge works", "app/core/evaluator.py · evaluate_answer"],
        ["How the dashboard is computed", "app/routers/analytics.py"],
        ["How CRUD is exposed", "app/routers/ingest.py"],
        ["How Redis is used", "app/telemetry.py · rate_limit / cache_get / cache_set"],
        ["How the UI uploads", "web/app/documents/page.tsx + web/app/api/documents/route.ts"],
        ["How the token stays server-side", "web/lib/server.ts · adminFetch"],
    ],
    [2.6, 4.3],
    font_size=8.6,
)

H2("A.3 Final advice")
P("Learn the flow, not every line. If you can explain why the API writes pending, why the worker exists, why the chunk contains its vector, and why hybrid search needs fusion, you can defend the project. Then practice the demo twice with a fresh document. The code is already written; your job in the viva is to show that you understand the design decisions behind it.")
P("Good luck.")

doc.save(str(OUT))
print(f"saved {OUT}")
