#!/usr/bin/env python3
"""Build the concise 11-slide Cogito deck: docs/Cogito_Presentation.pptx."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "Cogito_Presentation.pptx"

BG = RGBColor.from_string("F6F7FB")
PANEL = RGBColor.from_string("FFFFFF")
PANEL2 = RGBColor.from_string("EEF2F8")
LINE = RGBColor.from_string("D7DEEA")
TEXT = RGBColor.from_string("142033")
MUTED = RGBColor.from_string("5B6B82")
ACCENT = RGBColor.from_string("2563EB")
ACCENT2 = RGBColor.from_string("6D4AFF")
GOOD = RGBColor.from_string("0F9D6E")
WARN = RGBColor.from_string("B45309")
WHITE = RGBColor.from_string("FFFFFF")
CODEBG = RGBColor.from_string("0F172A")
FONT = "Arial"
MONO = "Consolas"
SW, SH = 13.333, 7.5
LEFT, CW = 0.65, 12.03

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background(); bg.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=PANEL, line=LINE, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    try: sh.adjustments[0] = radius
    except Exception: pass
    return sh


def text(s, x, y, w, h, value, size=13, color=TEXT, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.05, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for line in str(value).split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.line_spacing = line_spacing
        r = p.add_run(); r.text = line
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
        r.font.italic = italic; r.font.color.rgb = color
    return tb


def bullets(s, x, y, w, h, items, size=11.5, color=TEXT, bullet_color=ACCENT, gap=5):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(gap); p.line_spacing = 1.04
        lead = p.add_run(); lead.text = "•  "
        lead.font.name = FONT; lead.font.size = Pt(size); lead.font.color.rgb = bullet_color; lead.font.bold = True
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run(); r1.text = head; r1.font.name = FONT; r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = color
            r2 = p.add_run(); r2.text = rest; r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = color
        else:
            r = p.add_run(); r.text = item; r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def header(s, kicker, title, subtitle=None, num=None):
    text(s, LEFT, 0.24, 3.0, 0.25, "COGITO", 9, MUTED, bold=True)
    text(s, LEFT, 0.50, 9.0, 0.25, kicker.upper(), 9.5, ACCENT, bold=True)
    text(s, LEFT, 0.72, 11.8, 0.62, title, 27, TEXT, bold=True, line_spacing=1.0)
    if subtitle:
        text(s, LEFT, 1.33, 11.8, 0.30, subtitle, 11.5, MUTED)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(LEFT), Inches(1.64), Inches(CW), Inches(0.012))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background(); line.shadow.inherit = False
    if num is not None:
        text(s, 11.85, 7.14, 0.8, 0.2, f"{num:02d}", 8.5, MUTED, align=PP_ALIGN.RIGHT)
    text(s, LEFT, 7.14, 8.0, 0.2, "Cogito · MongoDB NoSQL project · October 2026", 8.5, MUTED)


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


def card(s, x, y, w, h, title, items, accent=ACCENT, body_size=11, title_size=13):
    box(s, x, y, w, h)
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.14), Inches(y + 0.16), Inches(0.05), Inches(0.3))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background(); bar.shadow.inherit = False
    text(s, x + 0.28, y + 0.12, w - 0.45, 0.28, title, title_size, TEXT, bold=True)
    if isinstance(items, str):
        text(s, x + 0.28, y + 0.50, w - 0.48, h - 0.62, items, body_size, MUTED, line_spacing=1.06)
    else:
        bullets(s, x + 0.28, y + 0.50, w - 0.48, h - 0.62, items, body_size, MUTED, gap=4)


def stat(s, x, y, w, h, label, value, sub=None, color=ACCENT):
    box(s, x, y, w, h)
    text(s, x + 0.16, y + 0.12, w - 0.28, 0.2, label.upper(), 8.5, MUTED, bold=True)
    text(s, x + 0.16, y + 0.33, w - 0.28, 0.42, value, 22, color, bold=True)
    if sub: text(s, x + 0.16, y + h - 0.38, w - 0.28, 0.26, sub, 8.5, MUTED)


def table(s, x, y, w, h, headers, rows, widths, font_size=9.0, header_size=9.0, highlight=None):
    shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table; tbl.first_row = True; tbl.horz_banding = False
    for i, cw in enumerate(widths): tbl.columns[i].width = Inches(cw)
    row_h = h / (len(rows) + 1)
    for i in range(len(rows) + 1): tbl.rows[i].height = Inches(row_h)
    for j, head in enumerate(headers):
        c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = PANEL2
        c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = head
        r.font.name = FONT; r.font.size = Pt(header_size); r.font.bold = True; r.font.color.rgb = ACCENT
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.fill.solid()
            c.fill.fore_color.rgb = PANEL if i % 2 else RGBColor.from_string("F7F9FC")
            if highlight == i: c.fill.fore_color.rgb = RGBColor.from_string("FFF6E5")
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; p.line_spacing = 0.96
            r = p.add_run(); r.text = str(val)
            r.font.name = FONT; r.font.size = Pt(font_size)
            r.font.color.rgb = TEXT if j == 0 else MUTED; r.font.bold = j == 0


def code(s, x, y, w, h, value, size=8.2):
    box(s, x, y, w, h, fill=CODEBG, line=LINE, radius=0.04)
    tb = s.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.12), Inches(w - 0.28), Inches(h - 0.2))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for line in value.strip("\n").split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(0); p.line_spacing = 1.0
        r = p.add_run(); r.text = line
        r.font.name = MONO; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string("DCE6F5")


def arrow(s, x, y, w, h, direction="right", label=None):
    if direction not in ("right", "left", "down", "up"):
        label, direction = direction, "right"
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW,
           "down": MSO_SHAPE.DOWN_ARROW, "up": MSO_SHAPE.UP_ARROW}[direction]
    sh = s.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = ACCENT2; sh.line.fill.background(); sh.shadow.inherit = False
    if label:
        if direction in ("right", "left"):
            text(s, x - 0.05, y - 0.22, w + 0.1, 0.18, label, 7.2, MUTED, align=PP_ALIGN.CENTER)
        else:
            text(s, x + w + 0.06, y, 1.0, 0.18, label, 7.2, MUTED)


def image_fit(s, path, x, y, w, h):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    nw, nh = iw * scale, ih * scale
    s.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2),
                         width=Inches(nw), height=Inches(nh))


# 1 ---------------------------------------------------------------- title
s = new_slide()
text(s, 8.6, 0.2, 4.4, 4.0, "∎", 230, RGBColor.from_string("E4EAF6"), bold=True)
text(s, LEFT, 0.75, 6.2, 0.3, "NOSQL DATABASES · COURSE PROJECT", 11, ACCENT, bold=True)
text(s, LEFT, 1.10, 7.0, 0.75, "Cogito", 54, TEXT, bold=True)
text(s, LEFT, 1.95, 7.2, 0.95, "A MongoDB-native document assistant with hybrid search and cited answers", 20, TEXT, line_spacing=1.1)
text(s, LEFT, 3.15, 7.0, 0.3, "One database. Five roles: store · search · vector · event bus · analytics.", 11.5, MUTED)
text(s, LEFT, 4.05, 7.0, 0.95,
     "Pranesh Nikhar · 2427010267\nSoumyajit Chatterjee · 2427010255\nAnmol Kaushik · 2427010576\nOjas Khandelwal · 2427010381",
     10.5, MUTED, line_spacing=1.35)
text(s, LEFT, 6.55, 6.0, 0.3, "October 2026", 10, MUTED)
for i, (title, sub, color) in enumerate([
    ("MongoDB", "documents · chunks · vectors", GOOD),
    ("mongot", "$search + $vectorSearch", ACCENT),
    ("LLM + embeddings", "pluggable · local fallback", ACCENT2),
]):
    y = 2.15 + i * 1.25
    box(s, 9.05, y, 3.55, 1.0)
    text(s, 9.25, y + 0.12, 3.15, 0.3, title, 15, color, bold=True)
    text(s, 9.25, y + 0.46, 3.15, 0.3, sub, 9.5, MUTED)
notes(s, "30-second intro: MongoDB is not just storage here — it is the search engine, vector store, event bus, and analytics engine.")


# 2 ---------------------------------------------------------------- criterion 1
s = new_slide()
header(s, "Criterion 1 · Problem identification and real-world relevance",
       "A real problem, a clear scope, and measurable objectives",
       "Evaluation objective: design, implement, and demonstrate a real-world NoSQL application.", 2)
card(s, LEFT, 1.82, 3.85, 3.55, "Evaluation objective", [
    "Design, implement, and demonstrate a real-world application using a NoSQL database.",
    "Assess practical implementation, data modeling, CRUD, querying, performance, and integration.",
    "Any suitable NoSQL database is allowed — we chose MongoDB.",
], ACCENT, body_size=9.8)
card(s, 4.74, 1.82, 3.85, 3.55, "Problem & real-world relevance", [
    "Knowledge is trapped in PDFs, contracts, policies, and notices.",
    "Keyword search misses paraphrases; vector-only search misses clause numbers.",
    "Unsourced chat answers create a trust gap.",
    "Users: legal, compliance, HR, operations, onboarding, audits.",
], ACCENT2, body_size=9.8)
card(s, 8.83, 1.82, 3.85, 3.55, "Project objectives", [
    "Zero-ETL ingestion through a MongoDB change stream.",
    "Hybrid retrieval with cited answers.",
    "Full document CRUD and aggregation analytics.",
    "One MongoDB deployment; no separate search or vector store.",
], GOOD, body_size=9.8)
box(s, LEFT, 5.55, CW, 1.0, fill=PANEL2, line=LINE)
text(s, LEFT + 0.2, 5.68, 2.0, 0.22, "SCOPE", 8.5, ACCENT, bold=True)
text(s, LEFT + 0.2, 5.92, 11.6, 0.48,
     "In scope: automatic ingestion, full CRUD, hybrid retrieval, cited Q&A, aggregation analytics, REST API.   "
     "Out of scope: multi-tenant login, HA cluster sizing, OCR quality, model fine-tuning.",
     9.4, MUTED, line_spacing=1.1)
text(s, LEFT, 6.72, CW, 0.25, "Success: upload → indexed → search → cited answer → edit → delete, with a refusal for an unrelated question.",
     9.0, MUTED)
notes(s, "State the course objective in one sentence, then the project problem. The scope strip keeps the demo honest: we are not claiming production multi-tenancy or HA. The success line is the flow the evaluator will see.")

# 3 ---------------------------------------------------------------- criterion 2
s = new_slide()
header(s, "Criterion 2 · NoSQL database selection and justification",
       "MongoDB fits the retrieval unit — and beats the alternatives here",
       "Justified against relational + pgvector and against a dedicated vector database.", 3)
card(s, LEFT, 1.82, 5.85, 2.25, "Why NoSQL", [
    "Variable document shapes: PDFs, contracts, notes, optional pages.",
    "Retrieval-first access pattern, not multi-table reporting.",
    "Text, vectors, and metadata are read together.",
], ACCENT, body_size=10.2)
card(s, 6.83, 1.82, 5.85, 2.25, "Why MongoDB", [
    "One JSON document = chunk + vector + metadata; no read-time join.",
    "Native $search + $vectorSearch + $rankFusion.",
    "Change streams, transactions, and aggregation in one system.",
], ACCENT2, body_size=10.2)
table(s, LEFT, 4.28, CW, 1.95, ["Option", "Fit for this workload", "Operational cost", "Verdict"], [
    ["MongoDB 8.2", "Store + keyword search + vector search + events + analytics", "One system, one backup", "Chosen"],
    ["Postgres + pgvector", "Strong SQL; retrieval needs joins + a separate lexical index", "More moving parts", "Viable alternative"],
    ["Dedicated vector DB", "Excellent ANN; weak metadata + no operational analytics", "Second store to sync", "Not chosen"],
], [1.55, 4.55, 3.6, 2.33], font_size=8.9, highlight=1)
text(s, LEFT, 6.38, CW, 0.55,
     "What we did: MongoDB 8.2 + mongot in Docker Compose; a chunk document holds text, vector, and metadata; no second store to sync. "
     "Relational wins when schema is fixed and multi-table reporting dominates — that is not this workload.",
     9.2, MUTED, line_spacing=1.08)
notes(s, "Give the one-sentence justification, then the comparison. The relational option is not wrong; it is a worse fit because the read is a chunk with its vector and metadata. Mention that MongoDB removes both a separate search engine and a vector database.")

# 4 ---------------------------------------------------------------- criterion 3
s = new_slide()
header(s, "Criterion 3 · Database design and data modeling",
       "Seven collections; chunks carry text, vector, and metadata",
       "Embedded vectors, optional fields, denormalized titles, access-pattern indexes.", 4)
table(s, LEFT, 1.82, 7.0, 4.15, ["Collection", "Holds", "Key fields"], [
    ["documents", "source + metadata", "path, title, status, raw"],
    ["chunks", "passages + vectors", "document_id, text, embedding"],
    ["conversations", "chat history", "user_id, messages[]"],
    ["memories", "long-term facts", "user_id, fact, embedding"],
    ["query_logs", "questions", "latency_ms, chunk_ids"],
    ["eval_results", "judge scores", "faithfulness, relevance"],
    ["feedback", "ratings (API only)", "answer_id, rating"],
], [1.35, 2.15, 3.5], font_size=8.7, header_size=8.7)
code(s, 7.9, 1.82, 4.78, 2.85, """
{
  "_id": "c392cfe9-...",
  "document_id": "82f9cb6d-...",
  "text": "4.1 Late Delivery. ...",
  "embedding": [0.013, -0.041, ...], // 768
  "page": 1, "order": 3,
  "token_count": 141,
  "title": "contract_penalties.md",
  "boost": 1.0
}
""", 7.8)
bullets(s, 7.9, 4.82, 4.78, 1.35, [
    "One document = text + vector + metadata; no join.",
    "Optional page/summary; title denormalized for citations.",
    "HNSW cosine index, 768d, scalar quantization.",
], 9.0, MUTED, gap=4)
text(s, LEFT, 6.18, CW, 0.62,
     "What we did: 7 collections; chunk as the retrieval aggregate; optional fields; denormalized title; text, search, HNSW, and compound indexes. "
     "Scalability: shard candidate document_id, read replicas, quantized vectors.",
     8.7, MUTED, line_spacing=1.08)
notes(s, "Point at the chunk JSON: text, vector, page, title, and boost in one document. That is the modeling argument. The bottom line answers the scalability and access-pattern part of the criterion.")

# 5 ---------------------------------------------------------------- criterion 4
s = new_slide()
header(s, "Criterion 4 · NoSQL database implementation and CRUD",
       "Create, read, update, delete — implemented and demonstrated",
       "Every operation is available in the API and the web UI; persistence is verified.", 5)
table(s, LEFT, 1.84, CW, 3.1, ["Operation", "API", "UI", "Correctness / persistence"], [
    ["Create", "POST /ingest, /ingest/raw", "Upload / paste", "10 MB cap, type whitelist, empty rejected"],
    ["Read", "GET /ingest/documents, /{id}, /chunks; POST /search", "Documents + Search", "404 handled; raw bytes hidden"],
    ["Update", "PATCH /ingest/documents/{id}; POST /{id}/reindex", "Rename / Save / Reindex", "Text edit → pending; chunk titles synced"],
    ["Delete", "DELETE /ingest/documents/{id}", "Delete + confirm", "Document + chunks in one transaction"],
], [1.35, 4.0, 2.55, 4.13], font_size=8.9)
card(s, LEFT, 5.28, 5.85, 1.45, "Persistence proof", [
    "MongoDB restart retains documents, chunks, and indexes.",
    "After delete, GET returns 404 and chunks are gone.",
], GOOD, body_size=9.5, title_size=12.5)
code(s, 6.83, 5.28, 5.85, 1.45, """
async with await client.start_session() as s:
    async with s.start_transaction():
        await chunks.delete_many({"document_id": doc_id}, session=s)
        await documents.delete_one({"_id": doc_id}, session=s)
""", 7.8)
notes(s, "Do not read the table. Say: create queues a pending document; update marks it pending and the change stream rebuilds chunks; delete removes the document and all chunks atomically. The UI exercises all four, including the live pending→indexed transition.")

# 6 ---------------------------------------------------------------- criterion 5
s = new_slide()
header(s, "Criterion 5 · Advanced NoSQL features and query optimization",
       "Change streams, hybrid search, aggregation — measured",
       "Features chosen for this workload, not a feature checklist.", 6)
card(s, LEFT, 1.82, 3.85, 3.2, "Event-driven", [
    "Insert pending → change stream fires.",
    "Worker extracts, chunks, embeds.",
    "Transaction: chunks + status flip.",
    "Resume token + startup backfill.",
], ACCENT, body_size=9.6, title_size=12.5)
card(s, 4.74, 1.82, 3.85, 3.2, "Hybrid retrieval", [
    "$search matches exact words and IDs.",
    "$vectorSearch matches meaning.",
    "$rankFusion / RRF: 1/(60+rank).",
    "Fallback: $text + cosine + RRF.",
], ACCENT2, body_size=9.6, title_size=12.5)
card(s, 8.83, 1.82, 3.85, 3.2, "Analytics + tuning", [
    "$facet, $bucketAuto, $group, $dateToString.",
    "Text, HNSW vector, and compound indexes.",
    "100 candidates for k=5 (20× recall margin).",
    "Optional cross-encoder reranker.",
], GOOD, body_size=9.6, title_size=12.5)
for i, (label, value, sub, color) in enumerate([
    ("Local ingest", "0.26 s", "first run", GOOD),
    ("Warm ingest", "6 ms", "local", GOOD),
    ("Hosted embed", "1.6–4.4 s", "network", WARN),
    ("Answer", "7.4 s", "chat + judge", WARN),
]):
    stat(s, LEFT + i * 3.06, 5.35, 2.85, 0.95, label, value, sub, color)
text(s, LEFT, 6.48, CW, 0.35,
     "What we did: replica set rs0 enables change streams + transactions; the vector branch over-requests candidates; the dashboard is aggregation-only.",
     9.0, MUTED)
notes(s, "Three pillars: change streams for ingestion, hybrid retrieval for search, aggregation for analytics. The 100-candidate ratio is the main ANN tuning knob. The measurements are from the verification run; hosted model calls are network-bound.")

# 7 ---------------------------------------------------------------- criterion 6
s = new_slide()
header(s, "Criterion 6 · Application integration and functionality",
       "Next.js + FastAPI + LLM, integrated end to end",
       "The UI is the product; one REST backend serves every client.", 7)
card(s, 0.85, 1.95, 2.6, 1.55, "Next.js 15", ["Documents · Chat · Dashboard", "CRUD + search + citations"], ACCENT, body_size=9.2, title_size=12.5)
card(s, 4.2, 1.95, 2.6, 1.55, "FastAPI backend", ["REST endpoints", "Async + Pydantic"], ACCENT2, body_size=9.2, title_size=12.5)
card(s, 8.05, 1.82, 3.55, 1.85, "MongoDB + mongot", ["documents · chunks · memories", "$search · $vectorSearch", "change streams · transactions"], GOOD, body_size=9.0, title_size=12.5)
card(s, 8.05, 4.55, 2.55, 1.2, "Ingest worker", ["change stream → chunks", "provider embeddings"], WARN, body_size=9.0, title_size=12.5)
card(s, 4.2, 4.55, 2.6, 1.2, "LLM + embeddings", ["chat · embed · judge", "provider-agnostic"], ACCENT2, body_size=9.2, title_size=12.5)
card(s, 0.85, 4.55, 2.6, 1.2, "Redis + Postgres", ["cache · rate limit", "analytics mirror"], ACCENT, body_size=9.2, title_size=12.5)
arrow(s, 3.50, 2.55, 0.55, 0.24, "REST")
arrow(s, 6.85, 2.55, 1.05, 0.24, "query")
arrow(s, 9.45, 3.80, 0.24, 0.55, "down")
arrow(s, 8.45, 3.80, 0.24, 0.55, "up")
arrow(s, 6.95, 5.02, 1.0, 0.24, "left")
arrow(s, 5.50, 3.62, 0.24, 0.78, "down")
arrow(s, 3.55, 5.02, 1.0, 0.24, "left")
box(s, LEFT, 6.02, 10.9, 0.62, fill=PANEL2, line=LINE)
text(s, 0.82, 6.12, 10.6, 0.42,
     "What we did: documents CRUD, cited chat, aggregation dashboard, server-side token proxy. "
     "Errors: 400 / 401 / 404 / 413 / 422 / 429 surfaced in the UI.",
     8.8, MUTED, line_spacing=1.05)
notes(s, "Trace a request: UI → FastAPI → MongoDB hybrid aggregation → LLM answer. The worker consumes the change stream separately. The token never reaches the browser.")

# 8 ---------------------------------------------------------------- criterion 7
s = new_slide()
header(s, "Criterion 7 · Testing, security, and data validation",
       "Tests, validation, and a protected write path",
       "Evidence is documented; limitations are stated honestly.", 8)
card(s, LEFT, 1.83, 6.4, 2.65, "Testing evidence", [
    "17 unit tests: chunking, embeddings, LLM, judge, rerank, extraction.",
    "CI: Ruff + pytest + npm run build.",
    "End-to-end API/UI verification.",
    "E2E: upload → indexed → search → ask → rename → delete.",
], ACCENT, body_size=9.7, title_size=12.5)
card(s, 7.25, 1.83, 5.43, 1.2, "Input validation", ["Pydantic lengths, k bounds, 10 MB cap, type whitelist."], ACCENT2, body_size=9.4, title_size=12)
card(s, 7.25, 3.13, 5.43, 1.2, "Authentication / access control", ["Admin token on writes; constant-time compare; server-side proxy."], ACCENT2, body_size=9.4, title_size=12)
card(s, 7.25, 4.43, 5.43, 1.2, "Rate limiting", ["60 requests/min per query prefix; 429 + Retry-After."], GOOD, body_size=9.4, title_size=12)
box(s, LEFT, 5.92, 6.4, 0.72, fill=RGBColor.from_string("FFF6E5"), line=RGBColor.from_string("F0C36B"))
text(s, LEFT + 0.16, 6.04, 6.1, 0.5, "Known gap: no automated Mongo integration suite in CI; the manual/Playwright E2E run is in TEST_REPORT.md.",
     8.7, WARN, line_spacing=1.04)
text(s, 7.25, 5.92, 5.43, 0.62, "pytest tests/ -q  →  17 passed\nruff check app  →  clean", 9.0, MUTED, font=MONO, line_spacing=1.2)
notes(s, "Be precise: 17 unit tests plus CI plus a manual E2E run. Do not overclaim an automated DB integration suite. The write path is protected; read endpoints remain open for the demo and that is listed as a limitation.")

# 9 ---------------------------------------------------------------- criterion 8
s = new_slide()
header(s, "Criterion 8 · Live demonstration, viva, and documentation",
       "A repeatable demo and a complete submission package",
       "The demo starts from an empty knowledge base and exercises the whole application.", 9)
steps = [
    ("0:00", "Empty knowledge base", "0 documents, 0 chunks"),
    ("0:30", "Upload a document", "pending → indexed"),
    ("1:30", "Search literal + paraphrase", "two different results"),
    ("3:00", "Ask a question", "citations shown"),
    ("4:30", "Ask an unrelated question", "refusal, no citations"),
    ("5:30", "Rename, edit, delete, dashboard", "full CRUD + analytics"),
]
for i, (t, title, sub) in enumerate(steps):
    y = 1.85 + i * 0.78
    text(s, LEFT, y + 0.04, 0.7, 0.22, t, 8.3, ACCENT, bold=True, font=MONO)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(LEFT + 0.78), Inches(y + 0.07), Inches(0.15), Inches(0.15))
    dot.fill.solid(); dot.fill.fore_color.rgb = ACCENT2; dot.line.fill.background(); dot.shadow.inherit = False
    text(s, LEFT + 1.02, y, 3.7, 0.26, title, 11.0, TEXT, bold=True)
    text(s, LEFT + 1.02, y + 0.25, 3.7, 0.22, sub, 8.5, MUTED)
card(s, 5.9, 1.82, 6.78, 2.55, "Documentation package", [
    "Report · feature map · viva sheet",
    "Demo script · test report · this deck",
], GOOD, body_size=10.0, title_size=12.5)
box(s, 5.9, 4.55, 6.78, 1.0, fill=RGBColor.from_string("FFF6E5"), line=RGBColor.from_string("F0C36B"))
text(s, 6.1, 4.70, 6.4, 0.7, "REQUIRED: plagiarism report (Turnitin/DrillBit). The repository cannot generate it.",
     10.0, WARN, line_spacing=1.05)
text(s, 5.9, 5.75, 6.78, 0.55, "Viva: 8 questions and answers in docs/VIVA_PREP.md. Demo/reset commands in docs/DEMO_SCRIPT.md.",
     9.0, MUTED, line_spacing=1.1)
text(s, LEFT, 6.65, 4.8, 0.25, "Start empty; add everything live.", 8.8, MUTED)
notes(s, "This is the demo and submission criterion. Follow the timeline; start empty and let the UI show pending→indexed, two different search results, a cited answer, a refusal, then CRUD and the dashboard. The plagiarism report is the one external artifact.")

# 10 --------------------------------------------------------------- rubric coverage
s = new_slide()
header(s, "Rubric coverage", "All eight criteria — what we did and where to prove it",
       "Use this as the submission map during the viva.", 10)
rows = [
    ["1", "Problem & relevance", "Document assistant; live upload; empty start", "Report §1–2; /documents"],
    ["2", "NoSQL selection", "MongoDB 8.2 + mongot; chunk aggregate; no second store", "Slide 3; DATABASE.md"],
    ["3", "Data modeling", "7 collections; embedded vectors; text/HNSW/compound indexes", "Slide 4; db.py"],
    ["4", "CRUD", "Create/read/update/delete; transactional delete", "Slide 5; ingest.py; UI"],
    ["5", "Advanced features", "Change streams, transactions, $search/$vectorSearch/$rankFusion, $facet", "Slide 6; hybrid_search.py"],
    ["6", "Integration", "Next.js + FastAPI + configurable LLM; 4xx handling", "Slide 7; web/"],
    ["7", "Testing & security", "17 tests, CI, E2E; Pydantic; token; 429", "Slide 8; TEST_REPORT.md"],
    ["8", "Demo, viva, docs", "Demo script, viva sheet, report, this deck", "Plagiarism report externally"],
]
table(s, LEFT, 1.84, CW, 4.45, ["#", "Criterion", "What we did", "Evidence / demo"], rows,
      [0.4, 2.15, 5.35, 4.13], font_size=8.7, header_size=9.0, highlight=8)
box(s, LEFT, 6.42, CW, 0.42, fill=RGBColor.from_string("FFF6E5"), line=RGBColor.from_string("F0C36B"))
text(s, LEFT + 0.16, 6.52, 11.6, 0.25,
     "Row 8: the plagiarism report is the one required artifact the repository cannot generate. Attach it before submission.",
     9.2, WARN)
notes(s, "This is the direct answer to “does it match the rubric?”. Every row names what we built and where to show it. The only external artifact is the plagiarism report.")

# 11 --------------------------------------------------------------- summary
s = new_slide()
header(s, "Summary", "What we delivered — and what is still honest work",
       "A complete NoSQL application with a clear boundary.", 11)
card(s, LEFT, 1.84, 6.2, 3.55, "Delivered", [
    "MongoDB store + mongot keyword/vector search",
    "Change-stream ingestion + transactional re-index",
    "Hybrid retrieval + cited answers",
    "Full CRUD API and web UI",
    "Aggregation dashboard",
    "REST API + configurable LLM/embeddings",
    "Pydantic validation, admin auth, 429 rate limit",
    "17 tests + CI + documented E2E",
], GOOD, body_size=9.6, title_size=13)
card(s, 7.1, 1.84, 5.58, 2.35, "Honest gaps / next", [
    "Per-user authentication and document ACL.",
    "Automated Mongo integration suite in CI.",
    "Production TLS, backups, and HA.",
], WARN, body_size=9.8, title_size=13)
box(s, 7.1, 4.42, 5.58, 0.97, fill=RGBColor.from_string("FFF6E5"), line=RGBColor.from_string("F0C36B"))
text(s, 7.3, 4.56, 5.2, 0.7, "REQUIRED: plagiarism report (external).", 9.8, WARN, line_spacing=1.05)
text(s, LEFT, 5.62, 6.2, 0.7, "The rubric is capability-based. We built a document assistant, a valid real-world NoSQL application, and mapped every criterion.",
     9.4, MUTED, line_spacing=1.08)
text(s, LEFT, 6.45, 11.8, 0.35, "One database. Five roles: store · search · vector · event bus · analytics.    Thank you.",
     12, MUTED, bold=True)
notes(s, "Close with what is actually delivered and the honest gaps. Then say the one-database line and invite questions. Do not claim production auth or HA.")

prs.save(str(OUT))
print(f"saved {OUT} ({len(prs.slides)} slides)")
